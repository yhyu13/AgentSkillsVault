#!/usr/bin/env python3
"""Generate examples/broken.xlsx — a tiny order sheet that trips every mechanical check.

Run:  python3 examples/make_sample.py
Then: python3 scripts/check_excel.py examples/broken.xlsx

Each planted error and the check it triggers:
  D3  "80" stored as text          -> Number as text
  F5  2.5 static among formulas      -> Number over formula
  F4  =D4/C4 with C4 = 0            -> Cell error (division by zero)
  B7  =SUM(B2:B5) but B6 populated  -> Truncated SUM
  E5  15 in a "%" column            -> Percent as whole number
"""
from pathlib import Path

import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Order"

# headers (row 1)
ws.append(["Item", "Qty", "Boxes", "Price", "Discount %", "Per box"])

# data (rows 2-6) — Item, Qty, Boxes, Price
rows = [
    ("Bolt M4",  1200, 24, 100),
    ("Washer",    800, 16, "80"),   # Price as TEXT -> Number as text (D3)
    ("Nut M4",    500,  0,  50),    # Boxes = 0     -> Cell error via Per box (F4)
    ("Screw",    1500, 30, 150),
    ("Bracket",   300,  6,  45),
]
for item, qty, boxes, price in rows:
    ws.append([item, qty, boxes, price])

# Discount % (col E) — fractions, except E5 typed as a whole number (should be 0.15)
for r, discount in zip(range(2, 7), [0.05, 0.10, 0.08, 15, 0.12]):
    ws.cell(r, 5, discount)

# Per box (col F) = Price / Boxes — formulas, except F5 overwritten with a constant
ws["F2"] = "=D2/C2"
ws["F3"] = "=D3/C3"
ws["F4"] = "=D4/C4"   # C4 = 0 -> division by zero
ws["F5"] = 2.5        # static number sitting among formulas
ws["F6"] = "=D6/C6"

# totals (row 7) — Qty SUM stops one row short
ws["A7"] = "Total"
ws["B7"] = "=SUM(B2:B5)"   # leaves out B6 -> Truncated SUM

out = Path(__file__).resolve().parent / "broken.xlsx"
wb.save(out)
print(f"wrote {out}")
