#!/usr/bin/env python3
"""excel-hygiene — generic MECHANICAL spreadsheet checks (no business rules).

Conservative, high-precision checks that apply to ANY xlsx. openpyxl-only (no LibreOffice).
Reads both formulas and numbers; recomputes from plain row numbers where needed. Assumes a simple
table with headers in row 1 of each sheet.

Checks:
  1. Number as text          - numeric value stored as a string -> SUM silently skips it.
  2. Number over formula     - a static number sitting in a column that is otherwise formulas.
  3. Cell error              - #DIV/0!/#REF!/... literals, and formulas that divide by zero.
  4. Truncated SUM           - a vertical SUM that stops before the last populated data row.
  5. Percent as whole number - a value > 1 in a column whose header contains '%'.

Usage:  python3 check_excel.py <file.xlsx>
"""
import sys
import re
import openpyxl

ERROR_LITERALS = {
    "#REF!", "#DIV/0!", "#N/A", "#VALUE!", "#NAME?", "#NULL!", "#NUM!",
    "#ССЫЛКА!", "#ДЕЛ/0!", "#Н/Д", "#ЗНАЧ!", "#ИМЯ?", "#ПУСТО!", "#ЧИСЛО!",
}


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_formula(cell):
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def looks_numeric(s):
    """A string that is really a number (with thousands spaces / nbsp / comma decimal)."""
    if not isinstance(s, str) or not s.strip():
        return False
    t = s.replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        float(t)
        return True
    except ValueError:
        return False


def resolve_numeric(ws, ref):
    """Numeric value of a cell ref ('N6'); follows one level of =SUM(a:b). None if not computable."""
    m = re.fullmatch(r"([A-Z]+)(\d+)", ref.strip())
    if not m:
        return None
    v = ws[ref.strip()].value
    if is_number(v):
        return float(v)
    if isinstance(v, str) and v.startswith("="):
        return resolve_sum(ws, v[1:])
    return None


def resolve_sum(ws, expr):
    """Sum the plain numbers in =SUM(a:b). None if the expr isn't a simple SUM."""
    m = re.fullmatch(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", expr.replace(" ", ""), re.I)
    if not m:
        return None
    total, seen = 0.0, False
    for row in ws[f"{m.group(1)}:{m.group(2)}"]:
        for c in row:
            if is_number(c.value):
                total += c.value
                seen = True
    return total if seen else None


def check_sheet(ws, findings):
    max_row, max_col = ws.max_row, ws.max_column

    # ---- per-column scan (checks 1, 2, 5) ----
    for col in range(1, max_col + 1):
        header = ws.cell(1, col).value
        header = header if isinstance(header, str) else ""
        is_percent_col = "%" in header
        n_num = n_formula = n_text = 0
        text_numeric, static_numbers, percent_outliers = [], [], []
        for row in range(2, max_row + 1):
            cell = ws.cell(row, col)
            if cell.value is None or cell.value == "":
                continue
            if is_formula(cell):
                n_formula += 1
            elif is_number(cell.value):
                n_num += 1
                static_numbers.append(cell)
                if is_percent_col and cell.value > 1:
                    percent_outliers.append(cell)
            elif isinstance(cell.value, str):
                if cell.value in ERROR_LITERALS:
                    continue  # handled in the formula/error pass
                n_text += 1
                if looks_numeric(cell.value):
                    text_numeric.append(cell)

        numeric_col = (n_num + n_formula) >= max(1, n_text)
        # 1) numbers stored as text (only in a column that is otherwise numeric)
        if numeric_col:
            for cell in text_numeric:
                findings.append(("Number as text", ws.title, cell.coordinate,
                                 f"value '{cell.value}' is stored as text -> SUM skips it, the total is understated"))
        # 2) a static number among formulas (formulas clearly dominant)
        if n_formula >= 2 and n_num >= 1 and n_formula >= 2 * n_num:
            for cell in static_numbers:
                findings.append(("Number over formula", ws.title, cell.coordinate,
                                 f"column '{header}' is formulas, but a static number {cell.value} sits here -- a formula may have been overwritten"))
        # 5) percent stored as a whole number
        for cell in percent_outliers:
            findings.append(("Percent as whole number", ws.title, cell.coordinate,
                             f"column '{header}' has {cell.value} > 1 -- likely should be {cell.value/100:g} (off by x100)"))

    # ---- per-cell formula scan (checks 3, 4) + error literals ----
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            v = cell.value
            # error literals (cached)
            if cell.data_type == "e" or (isinstance(v, str) and v in ERROR_LITERALS):
                findings.append(("Cell error", ws.title, cell.coordinate, f"cell contains the error {v}"))
                continue
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            formula = v[1:]
            # 3b) division by zero
            if "/" in formula:
                denom = formula.split("/")[-1].strip().strip("()")
                val = resolve_numeric(ws, denom)
                if val is None:
                    val = resolve_sum(ws, denom)
                if val is not None and val == 0:
                    findings.append(("Cell error", ws.title, cell.coordinate,
                                     f"division by zero ({denom}=0) -> #DIV/0!"))
            # 4) truncated vertical SUM (stops before the last populated data row)
            m = re.fullmatch(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", formula.replace(" ", ""), re.I)
            if m and m.group(1).upper() == m.group(3).upper():     # same column => vertical
                colL, r2 = m.group(1).upper(), int(m.group(4))
                nxt = ws[f"{colL}{r2 + 1}"]
                if is_number(nxt.value):                            # next row is real data, not the total
                    findings.append(("Truncated SUM", ws.title, cell.coordinate,
                                     f"SUM stops at row {r2}, but row {r2 + 1} is populated ({nxt.value}) -- a row may be dropped"))


def main():
    if len(sys.argv) < 2:
        print("usage: python3 check_excel.py <file.xlsx>")
        sys.exit(1)
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=False)
    findings = []
    for ws in wb.worksheets:
        check_sheet(ws, findings)

    print(f"\nExcel hygiene -- {path}")
    if not findings:
        print("  No mechanical errors found. (Business rules: separate skill.)")
        return
    print(f"  Mechanical findings: {len(findings)}\n")
    print(f"  {'class':<24} {'sheet':<12} {'cell':<7} what's wrong")
    print("  " + "-" * 90)
    for cls, sheet, coord, msg in findings:
        print(f"  {cls:<24} {sheet:<12} {coord:<7} {msg}")
    print("\n  These are mechanical checks (a second pair of eyes). Client-specific business rules belong in your own domain skill.")


if __name__ == "__main__":
    main()
